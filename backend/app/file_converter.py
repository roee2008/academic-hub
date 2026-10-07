import io
import os
import re
import html
import struct
import logging
from typing import Optional, List, Tuple, Dict, Any
from datetime import datetime

import olefile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    HRFlowable,
)
from PIL import Image

logger = logging.getLogger("academic_dashboard.file_converter")

# Formats natively recognized by Google NotebookLM
NOTEBOOKLM_NATIVE_EXTENSIONS = {
    ".pdf",
    ".pptx",
    ".docx",
    ".txt",
    ".md",
    ".mp3",
    ".wav",
    ".m4a",
    ".aac",
}

NOTEBOOKLM_NATIVE_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "text/markdown",
    "audio/mpeg",
    "audio/wav",
    "audio/mp4",
    "audio/x-m4a",
    "application/vnd.google-apps.document",
    "application/vnd.google-apps.presentation",
}

def get_file_extension(filename: str) -> str:
    _, ext = os.path.splitext(filename)
    return ext.lower()

def check_notebooklm_compatibility(title: str, mime_type: Optional[str] = None) -> Dict[str, Any]:
    """
    Evaluates whether a file can be directly ingested into NotebookLM
    or requires automated conversion/transfer.
    """
    ext = get_file_extension(title)
    mime = (mime_type or "").lower()

    if ext in NOTEBOOKLM_NATIVE_EXTENSIONS or mime in NOTEBOOKLM_NATIVE_MIMES:
        return {
            "is_compatible": True,
            "action": "DIRECT",
            "reason": "Natively supported by NotebookLM",
            "suggested_title": title,
            "target_mime": mime or "application/octet-stream",
            "original_extension": ext,
            "target_extension": ext,
        }

    # Legacy PowerPoint (.ppt)
    if ext == ".ppt" or "powerpoint" in mime or mime == "application/vnd.ms-powerpoint":
        base_name = os.path.splitext(title)[0]
        return {
            "is_compatible": False,
            "action": "CONVERT_PPT",
            "reason": "Legacy PowerPoint (.ppt) not supported by NotebookLM. Auto-transfers to .pptx & .pdf",
            "suggested_title": f"{base_name}.pptx",
            "target_mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "original_extension": ext,
            "target_extension": ".pptx",
        }

    # Legacy Word (.doc)
    if ext == ".doc" or "msword" in mime:
        base_name = os.path.splitext(title)[0]
        return {
            "is_compatible": False,
            "action": "CONVERT_DOC",
            "reason": "Legacy Word (.doc) not supported by NotebookLM. Auto-transfers to .docx & .pdf",
            "suggested_title": f"{base_name}.docx",
            "target_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "original_extension": ext,
            "target_extension": ".docx",
        }

    # Images
    if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"} or mime.startswith("image/"):
        base_name = os.path.splitext(title)[0]
        return {
            "is_compatible": False,
            "action": "CONVERT_IMAGE",
            "reason": "Images must be wrapped into PDF for NotebookLM source ingestion",
            "suggested_title": f"{base_name}.pdf",
            "target_mime": "application/pdf",
            "original_extension": ext,
            "target_extension": ".pdf",
        }

    # Spreadsheets
    if ext in {".xlsx", ".xls", ".csv"} or "spreadsheet" in mime or "excel" in mime or mime == "text/csv":
        base_name = os.path.splitext(title)[0]
        return {
            "is_compatible": False,
            "action": "CONVERT_SHEET",
            "reason": "Spreadsheets converted to structured Markdown/PDF tables for NotebookLM AI synthesis",
            "suggested_title": f"{base_name}_summary.pdf",
            "target_mime": "application/pdf",
            "original_extension": ext,
            "target_extension": ".pdf",
        }

    # Scripts / code / text-like
    if ext in {".py", ".asm", ".java", ".c", ".cpp", ".js", ".ts", ".html", ".css", ".json", ".xml", ".sql", ".sh"}:
        base_name = os.path.splitext(title)[0]
        return {
            "is_compatible": False,
            "action": "CONVERT_CODE",
            "reason": "Code files formatted as clean .txt / .pdf for NotebookLM ingestion",
            "suggested_title": f"{title}.txt",
            "target_mime": "text/plain",
            "original_extension": ext,
            "target_extension": ".txt",
        }

    # Generic fallback
    base_name = os.path.splitext(title)[0]
    return {
        "is_compatible": False,
        "action": "CONVERT_TEXT",
        "reason": "File converted to text/pdf format for NotebookLM readability",
        "suggested_title": f"{base_name}.txt",
        "target_mime": "text/plain",
        "original_extension": ext,
        "target_extension": ".txt",
    }

def _extract_strings_from_bytes(data: bytes, min_len: int = 4, max_items: int = 250) -> List[str]:
    """Extract printable strings from raw bytes efficiently with a set."""
    results = []
    seen = set()

    # ASCII / Latin-1 strings
    try:
        ascii_pattern = re.compile(b'[\x20-\x7e]{%d,}' % min_len)
        for match in ascii_pattern.finditer(data):
            if len(results) >= max_items:
                break
            try:
                s = match.group(0).decode('latin-1').strip()
                if len(s) >= min_len and not s.isdigit() and s not in seen:
                    seen.add(s)
                    results.append(s)
            except Exception:
                pass
    except Exception:
        pass

    # UTF-16LE strings (common in PPT binary format)
    try:
        utf16_pattern = re.compile(b'(?:[\x20-\x7e]\x00){%d,}' % min_len)
        for match in utf16_pattern.finditer(data):
            if len(results) >= max_items:
                break
            try:
                s = match.group(0).decode('utf-16le').strip()
                if len(s) >= min_len and not s.isdigit() and s not in seen:
                    seen.add(s)
                    results.append(s)
            except Exception:
                pass
    except Exception:
        pass

    return results

def extract_ppt_slides_data(ppt_bytes: bytes, original_title: str) -> List[Dict[str, Any]]:
    """
    Parses a legacy .ppt (PowerPoint 97-2003) file.
    Extracts slide structure, titles, and text content.
    """
    slides = []

    # 1. Try OLE file parsing
    try:
        if olefile.isOleFile(io.BytesIO(ppt_bytes)):
            with olefile.OleFileIO(io.BytesIO(ppt_bytes)) as ole:
                if ole.exists("PowerPoint Document"):
                    stream_data = ole.openstream("PowerPoint Document").read()
                    pos = 0
                    current_slide_texts = []
                    
                    while pos + 8 <= len(stream_data):
                        rec_ver_inst, rec_type, rec_len = struct.unpack('<HHI', stream_data[pos:pos+8])
                        pos += 8
                        is_container = (rec_ver_inst & 0x0F) == 0x0F

                        if is_container:
                            # Slide boundaries in PPT: SlideContainer (1006) or SlideListWithText (4080)
                            if rec_type in (1006, 4080):
                                if current_slide_texts:
                                    slides.append(current_slide_texts)
                                    current_slide_texts = []
                            # Container contents immediately follow; do not skip rec_len
                            continue

                        if pos + rec_len > len(stream_data):
                            break
                        rec_data = stream_data[pos:pos+rec_len]
                        pos += rec_len
                        
                        if rec_type == 4008:  # TextBytesAtom (1-byte characters)
                            text = rec_data.decode('latin-1', errors='replace').strip()
                            if text and len(text) > 1 and len(text) < 1000:
                                current_slide_texts.append(text)
                        elif rec_type == 4000:  # TextCharsAtom (2-byte UTF-16LE)
                            text = rec_data.decode('utf-16le', errors='replace').strip()
                            if text and len(text) > 1 and len(text) < 1000:
                                current_slide_texts.append(text)

                    if current_slide_texts:
                        slides.append(current_slide_texts)
    except Exception as e:
        logger.warning(f"OLE parsing error for {original_title}: {e}")

    # 2. Fallback if no slides extracted
    if not slides:
        extracted = _extract_strings_from_bytes(ppt_bytes)
        if not extracted:
            try:
                txt = ppt_bytes.decode('utf-8', errors='replace').strip()
                extracted = [line.strip() for line in txt.splitlines() if line.strip()]
            except Exception:
                extracted = [f"Presentation content for {original_title}"]

        # Group strings into simulated slides of 4-6 items
        chunk_size = 5
        for i in range(0, max(len(extracted), 1), chunk_size):
            chunk = extracted[i:i + chunk_size]
            if chunk:
                slides.append(chunk)

    # Clean and structure slides
    structured = []
    base_title = os.path.splitext(original_title)[0]

    for idx, slide_texts in enumerate(slides):
        filtered = [t for t in slide_texts if len(t) > 1]
        if not filtered:
            continue
        slide_title = filtered[0] if len(filtered[0]) < 80 else f"Slide {idx + 1}"
        bullets = filtered[1:] if len(filtered[0]) < 80 else filtered
        structured.append({
            "slide_number": idx + 1,
            "title": slide_title,
            "bullets": bullets,
        })

    if not structured:
        structured.append({
            "slide_number": 1,
            "title": base_title,
            "bullets": ["Slide content converted for Google NotebookLM reading."],
        })

    return structured

def convert_ppt_to_pptx(ppt_bytes: bytes, original_title: str) -> bytes:
    """
    Converts legacy .ppt content into a valid, native .pptx (PowerPoint OpenXML) presentation.
    """
    slides_data = extract_ppt_slides_data(ppt_bytes, original_title)
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # Title slide
    first_slide = prs.slides.add_slide(title_layout)
    base_title = os.path.splitext(original_title)[0]
    first_slide.shapes.title.text = base_title
    if first_slide.placeholders and len(first_slide.placeholders) > 1:
        first_slide.placeholders[1].text = f"Converted for Google NotebookLM\nOriginal File: {original_title}"

    # Content slides
    for s_info in slides_data:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = s_info["title"]
        
        # Add bullets to body placeholder
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        
        bullets = s_info["bullets"]
        if bullets:
            tf.text = bullets[0]
            for b in bullets[1:]:
                p = tf.add_paragraph()
                p.text = b
                p.level = 0
        else:
            tf.text = "Lecture notes & discussion points"

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out.getvalue()

def convert_ppt_to_pdf(ppt_bytes: bytes, original_title: str) -> bytes:
    """
    Converts legacy .ppt content into a clean, formatted PDF slide document.
    """
    slides_data = extract_ppt_slides_data(ppt_bytes, original_title)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=15,
    )
    slide_title_style = ParagraphStyle(
        'SlideTitle',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#4338CA'),
        spaceAfter=8,
    )
    bullet_style = ParagraphStyle(
        'SlideBullet',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        leftIndent=15,
        spaceAfter=4,
    )

    safe_doc_title = html.escape(original_title)
    story = [
        Paragraph(f"Presentation: {safe_doc_title}", title_style),
        Paragraph("<i>Auto-converted for Google NotebookLM AI synthesis</i>", styles['Italic']),
        Spacer(1, 15),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=15),
    ]

    for s_info in slides_data:
        try:
            safe_title = html.escape(str(s_info['title']))
            slide_elements = [
                Paragraph(f"Slide {s_info['slide_number']}: {safe_title}", slide_title_style),
                Spacer(1, 4),
            ]
            for b in s_info["bullets"]:
                clean_b = html.escape(str(b))
                slide_elements.append(Paragraph(f"• {clean_b}", bullet_style))
            slide_elements.append(Spacer(1, 12))
            slide_elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=12))
            story.append(KeepTogether(slide_elements))
        except Exception as e_s:
            logger.warning(f"Skipped malformed slide in PDF conversion: {e_s}")

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()

def convert_image_to_pdf(image_bytes: bytes, original_title: str) -> bytes:
    """Wraps an image into a clean PDF document for NotebookLM."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    story = [
        Paragraph(f"Visual Resource: {original_title}", styles['Heading1']),
        Spacer(1, 10),
    ]

    try:
        img = Image.open(io.BytesIO(image_bytes))
        w, h = img.size
        max_w, max_h = 500.0, 600.0
        ratio = min(max_w / w, max_h / h, 1.0)
        final_w, final_h = w * ratio, h * ratio

        from reportlab.platypus import Image as RLImage
        img_buf = io.BytesIO(image_bytes)
        story.append(RLImage(img_buf, width=final_w, height=final_h))
    except Exception as e:
        logger.warning(f"Could not format image into PDF: {e}")
        story.append(Paragraph(f"Image could not be rendered: {e}", styles['Normal']))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()

def convert_code_or_text_to_compatible(data_bytes: bytes, original_title: str) -> Tuple[bytes, str]:
    """Wraps code or raw text with a clean header for NotebookLM."""
    try:
        text = data_bytes.decode('utf-8', errors='replace')
    except Exception:
        text = data_bytes.decode('latin-1', errors='replace')

    header = f"=== File: {original_title} ===\n"
    header += f"=== Auto-prepared for Google NotebookLM ===\n\n"
    result = (header + text).encode('utf-8')
    return result, f"{original_title}.txt"

def generate_class_notebooklm_overview(
    course: Any,
    assignments: List[Any],
    timetable_slots: List[Any],
    files: List[Any],
) -> bytes:
    """
    Generates a master '00_NotebookLM_Course_Overview.pdf' providing
    high-level context (syllabus, timetable, assignments, document index)
    that grounds NotebookLM's AI.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'MainTitle',
        parent=styles['Title'],
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1E1B4B'),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#4F46E5'),
        spaceAfter=15,
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
    )

    story = [
        Paragraph(f"NotebookLM Knowledge Grounding: {course.name}", title_style),
        Paragraph(
            f"Course Code: <b>{course.code or 'N/A'}</b> | Section: <b>{course.section or 'General'}</b> | Generated: <b>{datetime.now().strftime('%b %d, %Y')}</b>",
            subtitle_style,
        ),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#6366F1'), spaceAfter=14),
    ]

    # Course Overview
    story.append(Paragraph("1. Course Description & Executive Summary", h2_style))
    summary_text = (
        f"This NotebookLM notebook consolidates all academic coursework, lecture materials, and syllabus "
        f"information for <b>{course.name}</b>. All documents included in this folder have been validated "
        f"and formatted for optimal AI grounding, audio overview generation, and question answering."
    )
    story.append(Paragraph(summary_text, body_style))
    story.append(Spacer(1, 10))

    # Weekly Bell Schedule
    if timetable_slots:
        story.append(Paragraph("2. Weekly Class Schedule", h2_style))
        day_names = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
        tt_data = [["Day", "Time", "Room Location"]]
        for slot in sorted(timetable_slots, key=lambda s: (s.day_of_week, s.start_time)):
            d_name = day_names[slot.day_of_week] if 0 <= slot.day_of_week <= 6 else f"Day {slot.day_of_week}"
            tt_data.append([d_name, f"{slot.start_time} - {slot.end_time}", slot.room or "Standard Classroom"])

        tt_table = Table(tt_data, colWidths=[120, 150, 200])
        tt_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF2FF')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#312E81')),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        story.append(tt_table)
        story.append(Spacer(1, 10))

    # Assignments & Homework
    if assignments:
        story.append(Paragraph("3. Assignments, Homework & Exams", h2_style))
        assign_data = [["Title", "Due Date", "Priority", "Status"]]
        for a in assignments[:15]:
            due = a.due_datetime.strftime('%b %d, %H:%M') if a.due_datetime else "No Deadline"
            assign_data.append([
                Paragraph(a.title[:45], body_style),
                due,
                a.priority or "MEDIUM",
                a.status or "TODO",
            ])

        a_table = Table(assign_data, colWidths=[200, 110, 80, 80])
        a_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        story.append(a_table)
        story.append(Spacer(1, 10))

    # Source Catalog & Conversion Status
    story.append(Paragraph("4. Catalog of Sources in this Notebook", h2_style))
    files_data = [["Document Title", "Original Type", "NotebookLM Status"]]
    for f in files:
        check = check_notebooklm_compatibility(f.title, f.mime_type)
        status_label = "Direct Ingestion" if check["is_compatible"] else f"Transferred ({check['target_extension']})"
        files_data.append([
            Paragraph(f.title[:50], body_style),
            check["original_extension"] or "Unknown",
            status_label,
        ])

    f_table = Table(files_data, colWidths=[250, 90, 130])
    f_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF2FF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#312E81')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(f_table)

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()
